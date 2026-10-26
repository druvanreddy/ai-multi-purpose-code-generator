"""
DMIXERAI
Document Extraction Service

Extracts plain text from uploaded documents for the AI Detector.
Supports: txt, md, html, docx, pdf.

Deliberately not supported: doc (legacy binary), rtf, odt — parsing these
reliably needs extra heavy/unreliable dependencies not worth adding for a
feature this scoped. Scanned/image-only PDFs are also not supported (no
OCR) — pypdf only reads embedded text.
"""

import io
import re

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {"txt", "md", "markdown", "html", "htm", "docx", "pdf"}


class ExtractionError(Exception):
    pass


def extract_text(filename, raw_bytes):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Unsupported file type '.{ext}'. Supported: "
            "txt, md, html, docx, pdf."
        )

    if ext in ("txt", "md", "markdown"):
        return _decode(raw_bytes)

    if ext in ("html", "htm"):
        return _strip_html(_decode(raw_bytes))

    if ext == "docx":
        return _extract_docx(raw_bytes)

    if ext == "pdf":
        return _extract_pdf(raw_bytes)

    raise ExtractionError(f"Unsupported file type '.{ext}'.")


def _decode(raw_bytes):
    try:
        return raw_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return raw_bytes.decode("latin-1", errors="ignore")


def _strip_html(html):
    html = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.DOTALL | re.IGNORECASE)
    html = re.sub(r"<[^>]+>", " ", html)
    html = re.sub(r"&nbsp;", " ", html)
    html = re.sub(r"&amp;", "&", html)
    html = re.sub(r"&lt;", "<", html)
    html = re.sub(r"&gt;", ">", html)
    html = re.sub(r"[ \t]+", " ", html)
    html = re.sub(r"\n{3,}", "\n\n", html)
    return html.strip()


def _extract_docx(raw_bytes):
    try:
        doc = Document(io.BytesIO(raw_bytes))
    except Exception as error:
        raise ExtractionError(f"Couldn't read this .docx file: {error}")

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts)

    if not text.strip():
        raise ExtractionError("No extractable text found in this .docx file.")

    return text


def _extract_pdf(raw_bytes):
    try:
        reader = PdfReader(io.BytesIO(raw_bytes))
    except Exception as error:
        raise ExtractionError(f"Couldn't read this PDF: {error}")

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ExtractionError("This PDF is password-protected and can't be read.")

    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue

    text = "\n\n".join(p.strip() for p in pages if p.strip())

    if not text.strip():
        raise ExtractionError(
            "No extractable text found — this may be a scanned/image-only "
            "PDF, which isn't supported (no OCR)."
        )

    return text
