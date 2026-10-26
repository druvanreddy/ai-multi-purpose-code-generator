"""
DMIXERAI
Report Export Service

Converts the report's own controlled Markdown (headings, paragraphs, bullet/
numbered lists, fenced code blocks, pipe tables, bold/italic/inline-code) into
DOCX (python-docx) and PDF (reportlab) documents.
"""

import io
import re

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


# =====================================================
# MARKDOWN PARSING (shared by both exporters)
# =====================================================

def _is_separator_row(line):
    s = line.strip()
    return bool(s) and "-" in s and all(c in "-|: " for c in s)


def parse_markdown_blocks(md):
    lines = md.replace("\r\n", "\n").split("\n")
    blocks = []
    i, n = 0, len(lines)

    while i < n:
        line = lines[i]

        if not line.strip():
            i += 1
            continue

        if line.strip().startswith("```"):
            lang = line.strip()[3:].strip()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "lang": lang, "text": "\n".join(code_lines)})
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            blocks.append({"type": f"h{len(m.group(1))}", "text": m.group(2).strip()})
            i += 1
            continue

        if "|" in line and i + 1 < n and _is_separator_row(lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = [header]
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append({"type": "table", "rows": rows})
            continue

        if re.match(r"^\s*[-*]\s+", line):
            items = []
            while i < n and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i]).strip())
                i += 1
            blocks.append({"type": "ul", "items": items})
            continue

        if re.match(r"^\s*\d+[.)]\s+", line):
            items = []
            while i < n and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i]).strip())
                i += 1
            blocks.append({"type": "ol", "items": items})
            continue

        para_lines = [line]
        i += 1
        while (
            i < n
            and lines[i].strip()
            and not re.match(r"^#{1,4}\s+", lines[i])
            and not lines[i].strip().startswith("```")
            and not re.match(r"^\s*[-*]\s+", lines[i])
            and not re.match(r"^\s*\d+[.)]\s+", lines[i])
        ):
            para_lines.append(lines[i])
            i += 1
        blocks.append({"type": "p", "text": " ".join(l.strip() for l in para_lines)})

    return blocks


# =====================================================
# DOCX EXPORT
# =====================================================

def _add_inline_runs_docx(paragraph, text):
    token_re = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`[^`]+`)")
    for part in token_re.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(part)


def _render_blocks_docx(doc, blocks):
    for b in blocks:
        t = b["type"]
        if t in ("h1", "h2", "h3", "h4"):
            doc.add_heading(b["text"], level=min(int(t[1]) + 1, 4))
        elif t == "p":
            _add_inline_runs_docx(doc.add_paragraph(), b["text"])
        elif t == "ul":
            for item in b["items"]:
                _add_inline_runs_docx(doc.add_paragraph(style="List Bullet"), item)
        elif t == "ol":
            for item in b["items"]:
                _add_inline_runs_docx(doc.add_paragraph(style="List Number"), item)
        elif t == "code":
            p = doc.add_paragraph()
            if b["lang"] == "mermaid":
                note = p.add_run("Mermaid diagram source (paste into a Mermaid renderer to view):")
                note.italic = True
                p = doc.add_paragraph()
            run = p.add_run(b["text"])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif t == "table" and b["rows"]:
            rows = b["rows"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    if c < len(table.rows[r].cells):
                        table.rows[r].cells[c].text = cell


def build_docx(title, chapters, references):
    doc = Document()
    doc.add_heading(title, level=0)

    for chapter in chapters:
        doc.add_heading(chapter["title"], level=1)
        _render_blocks_docx(doc, parse_markdown_blocks(chapter["content"]))

    if references:
        doc.add_heading("References", level=1)
        _render_blocks_docx(doc, parse_markdown_blocks(references))

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# =====================================================
# PDF EXPORT
# =====================================================

def _md_inline_to_rl(text):
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"\*(.+?)\*", r"<i>\1</i>", text)
    text = re.sub(r"`([^`]+)`", r"<font face='Courier'>\1</font>", text)
    return text


def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="H1Custom", parent=styles["Heading1"], spaceBefore=16, spaceAfter=10))
    styles.add(ParagraphStyle(name="H2Custom", parent=styles["Heading2"], spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle(name="H3Custom", parent=styles["Heading3"], spaceBefore=10, spaceAfter=6))
    styles.add(ParagraphStyle(name="BodyCustom", parent=styles["BodyText"], spaceAfter=8, leading=15))
    styles.add(ParagraphStyle(
        name="CodeCustom", parent=styles["Code"], backColor=colors.whitesmoke,
        borderPadding=6, fontSize=8.5, leading=11,
    ))
    return styles


def _render_blocks_pdf(md_text, styles):
    flow = []
    heading_style = {"h1": "H2Custom", "h2": "H2Custom", "h3": "H3Custom", "h4": "H3Custom"}

    for b in parse_markdown_blocks(md_text):
        t = b["type"]
        if t in heading_style:
            flow.append(Paragraph(_md_inline_to_rl(b["text"]), styles[heading_style[t]]))
        elif t == "p":
            flow.append(Paragraph(_md_inline_to_rl(b["text"]), styles["BodyCustom"]))
        elif t == "ul":
            items = [ListItem(Paragraph(_md_inline_to_rl(i), styles["BodyCustom"])) for i in b["items"]]
            flow.append(ListFlowable(items, bulletType="bullet"))
        elif t == "ol":
            items = [ListItem(Paragraph(_md_inline_to_rl(i), styles["BodyCustom"])) for i in b["items"]]
            flow.append(ListFlowable(items, bulletType="1"))
        elif t == "code":
            if b["lang"] == "mermaid":
                flow.append(Paragraph("<i>Mermaid diagram source (paste into a Mermaid renderer to view):</i>", styles["BodyCustom"]))
            safe = b["text"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            flow.append(Preformatted(safe, styles["CodeCustom"]))
            flow.append(Spacer(1, 6))
        elif t == "table" and b["rows"]:
            tbl = Table(b["rows"], hAlign="LEFT")
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efeaff")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]))
            flow.append(tbl)
            flow.append(Spacer(1, 10))

    return flow


def build_pdf(title, chapters, references):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
    )
    styles = _pdf_styles()

    story = [Paragraph(_md_inline_to_rl(title), styles["Title"]), Spacer(1, 20)]

    for idx, chapter in enumerate(chapters):
        if idx > 0:
            story.append(PageBreak())
        story.append(Paragraph(_md_inline_to_rl(chapter["title"]), styles["H1Custom"]))
        story.extend(_render_blocks_pdf(chapter["content"], styles))

    if references:
        story.append(PageBreak())
        story.append(Paragraph("References", styles["H1Custom"]))
        story.extend(_render_blocks_pdf(references, styles))

    doc.build(story)
    buffer.seek(0)
    return buffer
